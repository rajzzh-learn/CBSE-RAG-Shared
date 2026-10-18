"""
Shared file-attachment utilities for all CBSE Tutor projects.

Merges the best features from all 4 subject file_utils:
  - Chemistry: PDF, DOCX, TXT, images
  - CS / Maths / Physics: PDF, DOCX, PPTX, XLSX, HTML, TXT, images

Single entry point:
    text = extract_text(file_bytes, filename)
    uri, mime = image_to_base64_uri(file_bytes, filename)
    bool = is_image(filename)
"""
from __future__ import annotations

import base64
import io
from typing import NamedTuple


# ── MIME / extension maps ──────────────────────────────────────────────────

_IMAGE_MIME: dict[str, str] = {
    "png":  "image/png",
    "jpg":  "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
    "gif":  "image/gif",
}

SUPPORTED_IMAGE_EXTS: set[str] = set(_IMAGE_MIME)
SUPPORTED_DOC_EXTS: set[str]   = {"pdf", "docx", "pptx", "xlsx", "txt", "html", "htm",
                                   "md", "csv", "json", "tex"}
SUPPORTED_EXTS: set[str]       = SUPPORTED_IMAGE_EXTS | SUPPORTED_DOC_EXTS

DEFAULT_MAX_CHARS = 8000   # per-file text budget (chars)
TOTAL_TEXT_BUDGET = 12000  # combined budget when multiple files attached


class Attachment(NamedTuple):
    name: str
    is_image: bool
    text: str | None        # None for images
    image_uri: str | None   # None for documents
    preview: bytes | str    # bytes for images, str for docs


# ── Public helpers ─────────────────────────────────────────────────────────

def is_image(filename: str) -> bool:
    return _ext(filename) in SUPPORTED_IMAGE_EXTS


def image_to_base64_uri(file_bytes: bytes, filename: str) -> tuple[str, str]:
    ext  = _ext(filename)
    mime = _IMAGE_MIME.get(ext, "image/png")
    b64  = base64.b64encode(file_bytes).decode("ascii")
    return f"data:{mime};base64,{b64}", mime


def extract_text(file_bytes: bytes, filename: str,
                 max_chars: int = DEFAULT_MAX_CHARS) -> str:
    """Extract plain text from any supported document format."""
    ext = _ext(filename)
    if ext == "pdf":
        return _pdf(file_bytes, max_chars)
    if ext == "docx":
        return _docx(file_bytes, max_chars)
    if ext == "pptx":
        return _pptx(file_bytes, max_chars)
    if ext == "xlsx":
        return _xlsx(file_bytes, max_chars)
    if ext in ("html", "htm"):
        return _html(file_bytes, max_chars)
    if ext in ("txt", "md", "csv", "json", "tex"):
        return _txt(file_bytes, max_chars)
    raise ValueError(f"Unsupported file type: .{ext}")


def build_attachment(filename: str, file_bytes: bytes,
                     max_chars: int = DEFAULT_MAX_CHARS) -> Attachment:
    """Parse any uploaded file into a unified Attachment object."""
    if is_image(filename):
        uri, _ = image_to_base64_uri(file_bytes, filename)
        return Attachment(name=filename, is_image=True, text=None,
                          image_uri=uri, preview=file_bytes)
    text = extract_text(file_bytes, filename, max_chars)
    return Attachment(name=filename, is_image=False, text=text,
                      image_uri=None, preview=text)


# ── Format extractors ──────────────────────────────────────────────────────

def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _trim(text: str, max_chars: int) -> str:
    if len(text) > max_chars:
        return text[:max_chars] + f"\n\n[truncated at {max_chars} chars]"
    return text


def _pdf(file_bytes: bytes, max_chars: int) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [p.extract_text() or "" for p in reader.pages]
    return _trim("\n\n".join(t.strip() for t in pages if t.strip()), max_chars)


def _docx(file_bytes: bytes, max_chars: int) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _trim("\n".join(parts), max_chars)


def _pptx(file_bytes: bytes, max_chars: int) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    parts: list[str] = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        slide_texts.append(t)
        if slide_texts:
            parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_texts))
    return _trim("\n\n".join(parts), max_chars)


def _xlsx(file_bytes: bytes, max_chars: int) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
    return _trim("\n\n".join(parts), max_chars)


def _html(file_bytes: bytes, max_chars: int) -> str:
    try:
        import nh3
        raw = file_bytes.decode("utf-8", errors="replace")
        clean = nh3.clean(raw, tags=set(), attributes={})
        # Collapse whitespace
        import re
        clean = re.sub(r"\s{3,}", "\n\n", clean).strip()
        return _trim(clean, max_chars)
    except ImportError:
        # nh3 not installed — return raw text
        return _trim(file_bytes.decode("utf-8", errors="replace").strip(), max_chars)


def _txt(file_bytes: bytes, max_chars: int) -> str:
    return _trim(file_bytes.decode("utf-8", errors="replace").strip(), max_chars)
